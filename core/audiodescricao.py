"""Audiodescrição (P-06): espaços sem fala, verificação do roteiro, voz, faixas e exportações.

O roteiro é uma legenda comum (`.srt`) aberta no editor: cada "legenda" é uma descrição com início, fim e
texto. Este módulo não depende da interface.

- `espacos()`   trechos sem fala nem legenda onde cabe descrição, com quantas palavras cabem
- `verificar()` regras de tempo (cabe? cai em cima da fala? sobrepõe?) e de estilo do guia brasileiro
               (presente, sem interpretar, sem "vemos"), como avisos
- `Voz`         voz sintética do próprio sistema (Windows, Mac, Linux), para ouvir e para gerar a faixa
- `faixa_ad()`  faixa só com as descrições (wav), cada uma no seu tempo
- `mixar()`     áudio original com o volume abaixado nos trechos descritos + a faixa (wav estéreo)
- `roteiro_txt()`, `roteiro_docx()`, `ttml()` exportações do roteiro
- `quadros()` + `sugerir()` rascunho de descrição por IA a partir de quadros do vídeo (a pessoa revisa)
"""

from __future__ import annotations

import base64
import logging
import math
import os
import re
import subprocess
import sys
import wave
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np

log = logging.getLogger("legendas.audiodescricao")

PPM_PADRAO = 150            # palavras por minuto da voz da audiodescrição
MIN_ESPACO = 1.5            # s: menos que isso não cabe descrição útil
FOLGA = 0.2                 # s de respiro depois da fala e antes da próxima
_NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0


@dataclass
class Descricao:
    n: int
    inicio: float
    fim: float
    texto: str

    @property
    def palavras(self) -> int:
        return contar_palavras(self.texto)


def contar_palavras(texto: str) -> int:
    return len(re.findall(r"[\wÀ-ÿ]+(?:[-'][\wÀ-ÿ]+)*", texto or ""))


def palavras_cabem(duracao: float, ppm: float = PPM_PADRAO) -> int:
    return max(0, int(math.floor(duracao * ppm / 60.0 + 1e-6)))


def _juntar(trechos: list[tuple[float, float]]) -> list[tuple[float, float]]:
    out: list[list[float]] = []
    for a, b in sorted((a, b) for a, b in trechos if b > a):
        if out and a <= out[-1][1]:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return [(a, b) for a, b in out]


def falas_da_mascara(mask: np.ndarray, hop: float = 0.01) -> list[tuple[float, float]]:
    """Máscara 0/1 de voz (um valor a cada `hop` s) → trechos de fala (início, fim)."""
    m = np.concatenate([[0], np.asarray(mask, dtype=np.int8), [0]])
    d = np.flatnonzero(np.diff(m))
    return [(round(a * hop, 3), round(b * hop, 3)) for a, b in zip(d[0::2], d[1::2])]


def espacos(falas: list[tuple[float, float]], legendas: list[tuple[float, float]], duracao: float,
            minimo: float = MIN_ESPACO, folga: float = FOLGA, ppm: float = PPM_PADRAO) -> list[dict]:
    """Trechos livres (sem fala e sem legenda) com pelo menos `minimo` segundos depois das folgas."""
    ocupado = _juntar(list(falas) + list(legendas))
    livres, t = [], 0.0
    for a, b in ocupado + [(duracao, duracao)]:
        ini, fim = (t + folga if t > 0 else t), (a - folga if a < duracao else duracao)
        if fim - ini >= minimo:
            livres.append({"inicio": round(ini, 3), "fim": round(fim, 3), "duracao": round(fim - ini, 2),
                           "palavras": palavras_cabem(fim - ini, ppm)})
        t = max(t, b)
    return livres


def falas_do_audio(wav: str | Path, progresso=None) -> list[tuple[float, float]]:
    """Onde há fala no áudio (16 kHz mono). Usa o detector de voz; sem ele, os silêncios do ffmpeg."""
    from . import vad
    from .media import detect_silences, load_wav
    if vad.available():
        audio = load_wav(wav)
        return falas_da_mascara(vad.voice_mask(audio, progress=progresso))
    with wave.open(str(wav), "rb") as w:
        dur = w.getnframes() / w.getframerate()
    falas, t = [], 0.0
    for a, b in detect_silences(wav, -35.0, 0.4):
        if a > t:
            falas.append((t, a))
        t = b
    if t < dur:
        falas.append((t, dur))
    return falas


# ------------------------------------------------------------------ verificação

_PASSADO_EXCECOES = {"chapéu", "museu", "europeu", "europeus", "judeu", "adeus", "troféu", "mausoléu", "estou",
                     "ilhéu", "plebeu", "camafeu", "fogaréu", "escarcéu", "pigmeu", "liceu", "coliseu", "hebreu",
                     "macaréu", "mundaréu", "povaréu", "romeu", "tabaréu"}
_PASSADO_FORMAS = {"estava", "estavam", "tinha", "tinham", "era", "eram", "foi", "foram", "havia", "fez", "fizeram",
                   "disse", "pôs", "veio", "vieram", "deu", "viu", "saiu", "partiu", "caiu", "sorriu", "abriu", "subiu"}
_INTERPRETA = ["parece", "parecem", "aparenta", "aparentemente", "provavelmente", "talvez", "certamente", "obviamente",
               "claramente", "lindo", "linda", "lindos", "lindas", "bonito", "bonita", "feio", "feia", "horrível",
               "maravilhoso", "maravilhosa", "assustador", "assustadora", "feliz", "felizes", "triste", "tristes",
               "nervoso", "nervosa", "irritado", "irritada", "apaixonado", "apaixonada", "com medo", "com raiva"]
_VEMOS = re.compile(r"(?i)\b(vemos|vê-se|podemos ver|a câmera mostra|a cena mostra|é possível ver)\b")


def _sobrepoe(a: float, b: float, trechos: list[tuple[float, float]]) -> float:
    return sum(max(0.0, min(b, y) - max(a, x)) for x, y in trechos)


def verificar(descs: list[Descricao], falas: list[tuple[float, float]] | None = None, ppm: float = PPM_PADRAO,
              idioma: str = "pt-BR") -> list[dict]:
    """Apontamentos do roteiro: [{n, nivel: erro|aviso, regra, mensagem}]."""
    out: list[dict] = []
    ordem = sorted(descs, key=lambda d: d.inicio)
    pt = (idioma or "pt").lower().startswith("pt")

    def add(d, nivel, regra, msg):
        out.append({"n": d.n, "nivel": nivel, "regra": regra, "mensagem": msg})

    for i, d in enumerate(ordem):
        dur = d.fim - d.inicio
        if not d.texto.strip():
            add(d, "erro", "vazia", "Descrição sem texto.")
            continue
        cabe = palavras_cabem(dur, ppm)
        if d.palavras > cabe:
            add(d, "erro", "tempo", f"Não cabe: {d.palavras} palavras para {cabe} ({dur:.1f} s a {ppm:g} palavras/min).")
        if dur < 1.0:
            add(d, "aviso", "curta", "Trecho muito curto (menos de 1 segundo).")
        if i + 1 < len(ordem) and ordem[i + 1].inicio < d.fim - 0.01:
            add(d, "erro", "sobreposicao", f"Sobrepõe a descrição seguinte (nº {ordem[i + 1].n}).")
        if falas:
            s = _sobrepoe(d.inicio, d.fim, falas)
            if s > 0.3:
                add(d, "aviso", "fala", f"Cai em cima da fala ({s:.1f} s). Prefira os espaços sem diálogo.")
        if not pt:
            continue
        palavras = re.findall(r"[\wÀ-ÿ]+", d.texto.lower())
        passado = [w for w in palavras if w in _PASSADO_FORMAS or
                   (len(w) > 4 and re.search(r"(ou|eu|iu)$", w) and w not in _PASSADO_EXCECOES)]
        if passado:
            add(d, "aviso", "presente", f"Use o presente (“{passado[0]}” parece passado).")
        baixo = " " + " ".join(palavras) + " "
        interp = next((w for w in _INTERPRETA if f" {w} " in baixo), None)
        if interp:
            add(d, "aviso", "interpretar", f"Descreva o que se vê em vez de interpretar (“{interp}”).")
        if _VEMOS.search(d.texto):
            add(d, "aviso", "vemos", "Evite “vemos” / “a câmera mostra”: descreva direto.")
    return out


# ------------------------------------------------------------------ exportações do roteiro

def _tc(t: float, sep: str = ".") -> str:
    ms = int(round(max(0.0, t) * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


def roteiro_txt(descs: list[Descricao], titulo: str = "") -> str:
    linhas = [f"ROTEIRO DE AUDIODESCRIÇÃO{(' — ' + titulo) if titulo else ''}", ""]
    for d in sorted(descs, key=lambda d: d.inicio):
        linhas += [f"{d.n}  {_tc(d.inicio)} → {_tc(d.fim)}  ({d.fim - d.inicio:.1f} s)", d.texto.strip(), ""]
    return "\r\n".join(linhas)


def roteiro_docx(descs: list[Descricao], destino: str | Path, titulo: str = "") -> Path:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading("Roteiro de audiodescrição", level=1)
    if titulo:
        doc.add_paragraph(titulo)
    tab = doc.add_table(rows=1, cols=4)
    tab.style = "Table Grid"
    for c, t in zip(tab.rows[0].cells, ("Nº", "Entrada", "Saída", "Descrição")):
        c.text = t
    for d in sorted(descs, key=lambda d: d.inicio):
        row = tab.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = str(d.n), _tc(d.inicio), _tc(d.fim), d.texto.strip()
        for c in row:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    destino = Path(destino)
    doc.save(destino)
    return destino


def ttml(descs: list[Descricao], idioma: str = "pt-BR") -> str:
    """Texto com tempos (TTML) marcado como descrição, para plataformas que aceitam audiodescrição em texto."""
    corpo = "\n".join(
        f'      <p begin="{_tc(d.inicio)}" end="{_tc(d.fim)}">{escape(d.texto.strip()).replace(chr(10), "<br/>")}</p>'
        for d in sorted(descs, key=lambda d: d.inicio) if d.texto.strip())
    return ('<?xml version="1.0" encoding="UTF-8"?>\n'
            f'<tt xmlns="http://www.w3.org/ns/ttml" xmlns:ttm="http://www.w3.org/ns/ttml#metadata" xml:lang="{escape(idioma)}">\n'
            '  <head><metadata><ttm:title>Audiodescrição</ttm:title><ttm:desc>audio-description</ttm:desc></metadata></head>\n'
            '  <body>\n    <div ttm:role="description">\n' + corpo + '\n    </div>\n  </body>\n</tt>\n')


# ------------------------------------------------------------------ voz

class VozIndisponivel(RuntimeError):
    pass


class Voz:
    """Voz sintética do próprio sistema: Windows (voz do sistema), Mac (`say`), Linux (`espeak-ng`).

    `ppm` = palavras por minuto (a velocidade de leitura usada na verificação)."""

    def __init__(self, idioma: str = "pt-BR", ppm: float = PPM_PADRAO):
        self.idioma, self.ppm = idioma, ppm

    def falar(self, texto: str, destino: str | Path) -> Path:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        texto = " ".join(texto.split())
        if not texto:
            raise ValueError("Escreva a descrição antes de ouvir.")
        if os.environ.get("CNT_VOZ_FALSA") == "1":         # testes: um bipe com a duração da leitura
            dur = max(0.3, contar_palavras(texto) * 60.0 / self.ppm)
            t = np.arange(int(22050 * dur)) / 22050
            _grava_wav(destino, (0.2 * np.sin(2 * np.pi * 330 * t)), 22050)
            return destino
        if sys.platform == "win32":
            rate = max(-10, min(10, round((self.ppm - 150) / 12)))
            script = ("Add-Type -AssemblyName System.Speech; $s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
                      "$v = $s.GetInstalledVoices() | Where-Object { $_.VoiceInfo.Culture.Name -like $env:CNT_VOZ_LANG + '*' } | "
                      "Select-Object -First 1; if ($v) { $s.SelectVoice($v.VoiceInfo.Name) }; "
                      f"$s.Rate = {rate}; $s.SetOutputToWaveFile($env:CNT_VOZ_ARQ); $s.Speak($env:CNT_VOZ_TEXTO); $s.Dispose()")
            env = {**os.environ, "CNT_VOZ_TEXTO": texto, "CNT_VOZ_ARQ": str(destino), "CNT_VOZ_LANG": self.idioma.split("-")[0]}
            cmd = ["powershell", "-NoProfile", "-NonInteractive", "-Command", script]
        elif sys.platform == "darwin":
            env = None
            cmd = ["say", "-r", str(int(self.ppm)), "-o", str(destino), "--file-format=WAVE", "--data-format=LEI16@22050", texto]
        else:
            import shutil
            exe = shutil.which("espeak-ng") or shutil.which("espeak")
            if not exe:
                raise VozIndisponivel("Para ouvir a descrição no Linux, instale o leitor de voz \"espeak-ng\" "
                                      "(pela loja de programas do sistema).")
            env = None
            voz = "pt-br" if self.idioma.lower().startswith("pt") else self.idioma.split("-")[0]
            cmd = [exe, "-v", voz, "-s", str(int(self.ppm)), "-w", str(destino), texto]
        try:
            r = subprocess.run(cmd, capture_output=True, timeout=120, env=env, creationflags=_NO_WINDOW)
        except (OSError, subprocess.SubprocessError) as e:
            raise VozIndisponivel(f"Não consegui usar a voz do sistema: {e}") from e
        if r.returncode != 0 or not destino.exists() or destino.stat().st_size < 100:
            err = (r.stderr or b"").decode("utf-8", "replace").strip()[-200:]
            raise VozIndisponivel(f"A voz do sistema não gerou o áudio. {err}".strip())
        return destino


def _grava_wav(path: Path, x: np.ndarray, taxa: int) -> None:
    pcm = (np.clip(x, -1, 1) * 32767).astype(np.int16)
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(taxa)
        w.writeframes(pcm.tobytes())


def duracao_wav(path: str | Path) -> float:
    with wave.open(str(path), "rb") as w:
        return w.getnframes() / float(w.getframerate())


# ------------------------------------------------------------------ faixas

TAXA_SAIDA = 48000


def _para_48k(ffmpeg: str, src: Path, dst: Path) -> Path:
    subprocess.run([ffmpeg, "-y", "-nostdin", "-hide_banner", "-loglevel", "error", "-i", str(src),
                    "-ac", "1", "-ar", str(TAXA_SAIDA), "-c:a", "pcm_s16le", str(dst)],
                   check=True, capture_output=True, creationflags=_NO_WINDOW)
    return dst


def faixa_ad(falas_ad: list[tuple[float, Path]], duracao: float, destino: str | Path, ffmpeg: str) -> Path:
    """Uma faixa (wav 48 kHz mono) com cada áudio no seu tempo e silêncio no resto. Escreve aos poucos (filme
    inteiro não precisa caber na memória)."""
    destino = Path(destino)
    tmp = destino.with_suffix(".tmp48.wav")
    pos = 0                                    # amostras já escritas
    total = int(duracao * TAXA_SAIDA)
    with wave.open(str(destino), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(TAXA_SAIDA)
        for ini, arq in sorted(falas_ad, key=lambda x: x[0]):
            alvo = int(ini * TAXA_SAIDA)
            if alvo > pos:
                _silencio(w, alvo - pos)
                pos = alvo
            _para_48k(ffmpeg, Path(arq), tmp)
            with wave.open(str(tmp), "rb") as r:
                dados = r.readframes(r.getnframes())
            w.writeframes(dados)
            pos += len(dados) // 2
        if total > pos:
            _silencio(w, total - pos)
    tmp.unlink(missing_ok=True)
    return destino


def _silencio(w, n: int) -> None:
    bloco = b"\0\0" * TAXA_SAIDA
    while n > 0:
        k = min(n, TAXA_SAIDA)
        w.writeframes(bloco[: k * 2])
        n -= k


def expressao_volume(trechos: list[tuple[float, float]], abaixar: float = 0.3, rampa: float = 0.3) -> str:
    """Volume do áudio original: 1 fora das descrições, `abaixar` durante (com rampa suave antes e depois)."""
    if not trechos:
        return "1"
    partes = [f"clip(min((t-{a - rampa:.3f})/{rampa},({b + rampa:.3f}-t)/{rampa}),0,1)" for a, b in _juntar(trechos)]
    soma = "+".join(partes)
    return f"1-{1 - abaixar:.3f}*clip({soma},0,1)"


def mixar(video: str | Path, faixa: str | Path, trechos: list[tuple[float, float]], destino: str | Path, ffmpeg: str,
          abaixar: float = 0.3, pasta_tmp: str | Path | None = None) -> Path:
    """Áudio original (estéreo) com o volume abaixado nos trechos descritos + a faixa de audiodescrição."""
    destino = Path(destino)
    script = Path(pasta_tmp or destino.parent) / "mix_filtro.txt"
    script.write_text(
        f"[0:a:0]aformat=channel_layouts=stereo,volume='{expressao_volume(trechos, abaixar)}':eval=frame[base];"
        "[1:a]aformat=channel_layouts=stereo,volume=1.0[ad];"
        "[base][ad]amix=inputs=2:normalize=0:duration=first[mix]", encoding="utf-8")
    r = subprocess.run([ffmpeg, "-y", "-nostdin", "-hide_banner", "-loglevel", "error", "-i", str(video), "-i", str(faixa),
                        "-filter_complex_script", str(script), "-map", "[mix]", "-ar", str(TAXA_SAIDA),
                        "-c:a", "pcm_s16le", str(destino)], capture_output=True, creationflags=_NO_WINDOW)
    script.unlink(missing_ok=True)
    if r.returncode != 0:
        raise RuntimeError("Não consegui mixar o áudio: " + (r.stderr or b"").decode("utf-8", "replace").strip()[-300:])
    return destino


# ------------------------------------------------------------------ sugestão por IA

def quadros(video: str | Path, inicio: float, fim: float, pasta: str | Path, ffmpeg: str, n: int = 3) -> list[Path]:
    """`n` quadros (jpg, 512 px de largura) espalhados pelo trecho."""
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    out = []
    for i in range(n):
        t = inicio + (fim - inicio) * (i + 0.5) / n
        f = pasta / f"q{i}.jpg"
        r = subprocess.run([ffmpeg, "-y", "-nostdin", "-hide_banner", "-loglevel", "error", "-ss", f"{max(0.0, t):.3f}",
                            "-i", str(video), "-frames:v", "1", "-vf", "scale=512:-2", "-q:v", "4", str(f)],
                           capture_output=True, creationflags=_NO_WINDOW)
        if r.returncode == 0 and f.exists():
            out.append(f)
    return out


INSTRUCOES = (
    "Você escreve audiodescrição em português do Brasil para pessoas cegas, seguindo o guia brasileiro: "
    "descreva só o que se vê nos quadros (ações, pessoas, lugar, textos na tela), no presente, em frases curtas, "
    "sem interpretar sentimentos nem intenções, sem antecipar o que vai acontecer e sem dizer “vemos”. "
    "Use no máximo {palavras} palavras. Responda só com JSON: {{\"itens\": [{{\"id\": 1, \"texto\": \"...\"}}]}}")


def conteudo_com_imagens(tradutor, texto: str, imagens: list[Path]):
    """Mensagem do usuário com imagens, no formato que o serviço de IA escolhido entende."""
    from .translate import AnthropicTranslator
    b64 = [base64.b64encode(Path(p).read_bytes()).decode("ascii") for p in imagens]
    if isinstance(tradutor, AnthropicTranslator):
        return [*({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b}} for b in b64),
                {"type": "text", "text": texto}]
    return [{"type": "text", "text": texto},
            *({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b}"}} for b in b64)]


def sugerir(tradutor, imagens: list[Path], palavras: int, contexto: str = "") -> str:
    """Rascunho de descrição a partir dos quadros. Nada é aplicado sem a pessoa aceitar."""
    if not imagens:
        raise ValueError("Não consegui tirar quadros do vídeo nesse trecho.")
    pedido = (f"Quadros do trecho a descrever (em ordem). Cabem até {palavras} palavras."
              + (f"\nDescrição anterior (para não repetir): {contexto}" if contexto else ""))
    dados, _ = tradutor.translate_batch(INSTRUCOES.format(palavras=max(3, palavras)),
                                        conteudo_com_imagens(tradutor, pedido, imagens))
    itens = dados.get("itens") or []
    texto = str(itens[0].get("texto", "")).strip() if itens else ""
    if not texto:
        raise ValueError("A IA não devolveu uma descrição. Tente de novo.")
    return texto
