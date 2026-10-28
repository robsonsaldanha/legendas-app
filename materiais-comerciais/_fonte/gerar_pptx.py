from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
IMG='../img/'
BG=RGBColor(0x16,0x19,0x1D); PANEL=RGBColor(0x1F,0x23,0x29); P2=RGBColor(0x26,0x2B,0x32); LINE=RGBColor(0x3A,0x41,0x4B)
ACC=RGBColor(0xDD,0xB0,0x6A); ACC2=RGBColor(0xEB,0xC5,0x87); INK=RGBColor(0xE9,0xEC,0xEF); MUT=RGBColor(0xA3,0xAD,0xB8); ON=RGBColor(0x1F,0x16,0x06)
HEAD='Impact'; BODY='Calibri'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
W,H=13.333,7.5
def new(notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s
def box(s,x,y,w,h,fill=PANEL,line=LINE,r=0.05):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.adjustments[0]=r
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(1)
    sh.shadow.inherit=False; return sh
def txt(s,x,y,w,h,text,size=16,color=INK,font=BODY,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(0.02); tf.margin_top=tf.margin_bottom=Inches(0.02)
    lines=text if isinstance(text,list) else [text]
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        if spacing: p.line_spacing=spacing
        r=p.add_run(); r.text=l; r.font.size=Pt(size); r.font.color.rgb=color; r.font.name=font; r.font.bold=bold
    return tb
def rich(s,x,y,w,h,parts,size=16,align=PP_ALIGN.LEFT):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    for t,c,f in parts:
        r=p.add_run(); r.text=t; r.font.size=Pt(size); r.font.color.rgb=c; r.font.name=f
    return tb
def title(s,eyebrow,t,y=0.55):
    txt(s,0.7,y,8,0.3,eyebrow.upper(),11,ACC,BODY,True)
    txt(s,0.7,y+0.3,11.8,1.1,t,36,INK,HEAD)
def pic(s,name,x,y,w):
    p=s.shapes.add_picture(IMG+name,Inches(x),Inches(y),width=Inches(w)); p.line.color.rgb=LINE; p.line.width=Pt(1); return p
def footer(s,n):
    ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(7.12),prs.slide_width,Inches(0.38)); ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor(0x1A,0x1E,0x23); ln.line.fill.background()
    txt(s,0.7,7.17,6,0.25,'CNT Legendas  ·  legendas.cntoficial.com.br',10,MUT)
    txt(s,6.9,7.17,5.7,0.25,f'Aplicativo em constante atualização  ·  {n}',10,MUT,align=PP_ALIGN.RIGHT)
def tag(s,x,y,t,w=3.2):
    b=box(s,x,y,w,0.38,RGBColor(0x3A,0x32,0x24),ACC,0.5); txt(s,x,y+0.04,w,0.3,t,12,ACC2,BODY,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
def bullets(s,x,y,w,items,size=15,gap=0.62):
    for i,(h,d) in enumerate(items):
        yy=y+i*gap
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(yy+0.06),Inches(0.2),Inches(0.2)); c.fill.solid(); c.fill.fore_color.rgb=ACC; c.line.fill.background()
        rich(s,x+0.35,yy,w-0.35,gap,[(h+'  ',INK,BODY),(d,MUT,BODY)],size)
def card(s,x,y,w,h,head,body,num=None):
    box(s,x,y,w,h)
    if num: txt(s,x+0.25,y+0.18,1,0.6,num,30,ACC,HEAD)
    off=0.8 if num else 0.2
    txt(s,x+0.25,y+off,w-0.5,0.4,head,16,INK,BODY,True)
    txt(s,x+0.25,y+off+0.42,w-0.5,h-off-0.5,body,12.5,MUT,BODY)
n=[0]
def nxt(): n[0]+=1; return n[0]
def logo(s,x=0.7,y=0.5,w=1.7): s.shapes.add_picture(IMG+'logo.png',Inches(x),Inches(y),width=Inches(w))

# 1 capa
s=new('Abertura: apresente o CNT Legendas em uma frase. "Programa de legendagem profissional que traduz, revisa, sincroniza e confere as legendas, rodando no computador da equipe." Destaque que é um aplicativo em constante atualização.')
logo(s); txt(s,0.7,2.0,6.5,0.3,'LEGENDAGEM PROFISSIONAL',12,ACC,BODY,True)
txt(s,0.7,2.4,6.3,2.4,['Do arquivo do','cliente à entrega,','num só lugar.'],46,INK,HEAD,spacing=0.95)
txt(s,0.7,4.95,5.8,1.2,'Traduza, revise, sincronize e confira suas legendas no padrão das grandes plataformas, com a velocidade da inteligência artificial e o controle da sua equipe.',16,MUT)
tag(s,0.7,6.15,'Aplicativo em constante atualização',3.7); tag(s,4.55,6.15,'Avaliação gratuita de 24 horas',3.2)
pic(s,'editor.jpg',7.1,1.7,5.7)
nxt()

# 2 desafio
s=new('Conecte com a dor do cliente: prazos curtos, manuais diferentes por cliente e retrabalho. Pergunte quantos episódios por mês eles legendam e quanto tempo gastam na revisão.')
title(s,'O desafio','Legendar bem, rápido e\ndo jeito de cada cliente.')
for i,(h,d) in enumerate([('Prazo apertado','Episódios chegam em lote e o prazo não muda. Cada hora em tarefa repetitiva é uma hora a menos de revisão.'),('Regras de cada cliente','Limite de caracteres, velocidade de leitura, pontuação, glossário de nomes. O que vale para um não vale para outro.'),('Qualidade sem retrabalho','Erro de tempo ou de estilo achado na entrega vira devolução. O ideal é pegar tudo antes de enviar.')]):
    card(s,0.7+i*4.1,2.65,3.85,2.5,h,d,str(i+1))
box(s,0.7,5.45,12,1.25,RGBColor(0x2B,0x26,0x1C),ACC)
txt(s,1.0,5.6,3.3,1.0,['A resposta:','CNT Legendas'],26,ACC2,HEAD,anchor=MSO_ANCHOR.MIDDLE)
txt(s,4.5,5.6,8,1.0,'Um programa único: tradução com IA, editor profissional, sincronia pela fala, verificação automática e entrega. Instala em minutos, roda no seu computador e já vem com os padrões de legendagem prontos.',15,INK,anchor=MSO_ANCHOR.MIDDLE)
footer(s,2)

# 3 como funciona
s=new('Mostre o fluxo em quatro passos. Reforce: os arquivos originais do cliente nunca são alterados.')
title(s,'Como funciona','Quatro passos, do início ao fim.')
for i,(h,d) in enumerate([('Arquivos','Escolha a legenda do cliente e o vídeo. Os originais ficam intactos.'),('Traduzir','A IA traduz em lotes, com contexto, glossário e regras do cliente.'),('Revisar','Editor com vídeo, forma de onda e conferência automática.'),('Entregar','Arquivo no formato e no nome que o cliente pede.')]):
    card(s,0.7+(i%2)*3.15,1.95+(i//2)*2.4,3.0,2.25,h,d,f'0{i+1}')
pic(s,'projeto.jpg',7.1,2.0,5.7)
footer(s,3)

# 4 editor
s=new('Mostre o editor: grade com velocidade de leitura em cores, vídeo com a legenda, forma de onda arrastável. Tudo em uma tela.')
txt(s,0.7,0.8,4,0.3,'EDITOR PROFISSIONAL',11,ACC,BODY,True); txt(s,0.7,1.1,4.2,1.6,['Tudo o que a','revisão precisa,','numa tela.'],34,INK,HEAD,spacing=0.95)
bullets(s,0.7,3.0,4.0,[('Grade','com velocidade de leitura em cores'),('Vídeo','com a legenda sobreposta, quadro exato'),('Forma de onda','para ajustar entradas e saídas'),('Desfazer e histórico','de versões, localizar, corretor')],14,0.85)
pic(s,'editor.jpg',5.0,1.35,7.9); footer(s,4)

# 5 tradução
s=new('Explique os dois jeitos de traduzir: pela API (automático, com custo do próprio cliente) ou pelo assistente de IA do cliente (sem API). Glossário e regras do cliente entram nos dois.')
pic(s,'traduzir.jpg',0.7,1.35,7.4)
txt(s,8.5,0.8,4.2,0.3,'TRADUÇÃO COM IA',11,ACC,BODY,True); txt(s,8.5,1.1,4.4,1.6,['Dois jeitos de','traduzir, um','resultado.'],34,INK,HEAD,spacing=0.95)
bullets(s,8.5,3.0,4.2,[('Automática','lotes com contexto, glossário e limite de caracteres'),('Sem API','pelo seu assistente de IA, com copiar e colar guiado'),('Revisão','ortografia, gramática e velocidade de leitura'),('Seguro','a origem nunca é alterada')],13.5,0.95)
footer(s,5)

# 6 sincronia
s=new('Sincronia é um diferencial: a legenda é acertada pela voz do vídeo. Há também sincronização por referência e por cortes de cena.')
txt(s,0.7,0.8,4,0.3,'SINCRONIA',11,ACC,BODY,True); txt(s,0.7,1.1,4.2,1.6,['Tempo certo,','na fala certa.'],34,INK,HEAD,spacing=0.95)
bullets(s,0.7,2.6,4.0,[('Pela fala','separa a voz da música e do ruído'),('Por referência','outra legenda, mesmo em outro idioma'),('Ajustes','atrasar, dois pontos, quadros por segundo'),('Cortes de cena','detectados e respeitados'),('Roteiro','texto sem tempo vira legenda')],13.5,0.78)
pic(s,'sincronizar.jpg',5.0,1.35,7.9); footer(s,6)

# 7 qualidade
s=new('Verificar antes de entregar reduz devolução. As regras vêm do perfil do cliente.')
pic(s,'verificar.jpg',0.7,1.35,7.4)
txt(s,8.5,0.8,4.2,0.3,'QUALIDADE E ENTREGA',11,ACC,BODY,True); txt(s,8.5,1.1,4.4,1.6,['Confira antes','de enviar.'],34,INK,HEAD,spacing=0.95)
bullets(s,8.5,2.6,4.2,[('Verificar','leitura, duração, intervalo, linhas'),('Corrigir','ajuste automático de tempos e estilo'),('Entregar','.srt, .vtt, .ttml ou .ass'),('Relatório','arquivo de verificação opcional')],13.5,0.9)
footer(s,7)

# 8 equipes
s=new('Para empresas: relatórios de produção, clientes e glossários, Workspace da empresa nos planos Business e Enterprise.')
txt(s,0.7,0.8,5,0.3,'PARA EQUIPES E EMPRESAS',11,ACC,BODY,True); txt(s,0.7,1.1,4.6,1.8,['Visão de quem','produz e de quem','gerencia.'],34,INK,HEAD,spacing=0.95)
bullets(s,0.7,3.1,4.4,[('Clientes','com glossário e regras próprios'),('Glossário','geral, do cliente e do projeto'),('Relatórios','produção, tempo e produtividade'),('Prazos','de entrega por episódio'),('Workspace','da empresa (Business e Enterprise)')],13.5,0.75)
pic(s,'relatorios.jpg',5.4,1.35,7.5); footer(s,8)

# 9 além da legenda
s=new('Mostre que o mesmo programa cobre legenda ao vivo e audiodescrição, e transcreve vídeo sem legenda.')
title(s,'Além da legenda','Acessibilidade no mesmo programa.')
for i,(h,d) in enumerate([('Legenda ao vivo','Gera a legenda na hora, com fila para correção, saída para o programa de transmissão e tradução ao vivo.'),('Audiodescrição','Acha espaços sem fala, sugere descrições com IA, ouve com voz do sistema ou gravada, e exporta roteiro e áudio.'),('Vídeo sem legenda?','Transcrição no computador, sem custo, com detecção automática do idioma.')]):
    card(s,0.7+i*4.1,1.9,3.85,1.95,h,d)
pic(s,'ao-vivo.jpg',4.37,4.05,4.6)
footer(s,9)

# 10 segurança
s=new('Reforce que o trabalho é local: relevante para empresas com conteúdo confidencial.')
title(s,'Segurança e privacidade','Seu trabalho é seu.')
for i,(h,d) in enumerate([('Roda no seu computador','Vídeos e legendas ficam na sua máquina. Na tradução com IA, só o texto das legendas vai ao serviço que você escolher.'),('Chaves protegidas','A chave da IA fica no cofre de senhas do computador, nunca em arquivo de texto.'),('Originais intocados','O programa nunca altera nem apaga seus arquivos. Resultados saem ao lado do original.'),('Licença e LGPD','Licença assinada, ligada ao computador. Termos e política de privacidade publicados.')]):
    card(s,0.7+(i%2)*6.1,1.95+(i//2)*2.35,5.9,2.1,h,d)
footer(s,10)

# 11 atualização
s=new('IMPORTANTE: destaque que é um aplicativo em constante atualização, com atualização automática ao abrir. Cite exemplos recentes (legenda ao vivo, audiodescrição, Mac e Linux).')
tag(s,0.7,0.65,'Em constante atualização',3.0); txt(s,0.7,1.2,6.2,1.9,['Um aplicativo que','não para de evoluir.'],36,INK,HEAD,spacing=0.95)
txt(s,0.7,2.9,5.8,1.6,'O CNT Legendas é desenvolvido em contato direto com quem legenda no dia a dia. Ao abrir, ele procura a versão mais nova e se atualiza sozinho, sem instalador e sem perder seu trabalho, mostrando o que mudou.',15,MUT)
for i,(a,b) in enumerate([('2.24','versão atual'),('20','idiomas'),('3','sistemas')]):
    box(s,0.7+i*2.05,4.7,1.9,1.3); txt(s,0.7+i*2.05,4.78,1.9,0.7,a,30,ACC,HEAD,align=PP_ALIGN.CENTER); txt(s,0.7+i*2.05,5.5,1.9,0.4,b,12,MUT,align=PP_ALIGN.CENTER)
txt(s,0.7,6.2,6,0.6,'Windows, Mac e Linux.',12,MUT)
txt(s,7.4,0.85,5,0.3,'EVOLUÇÕES RECENTES',11,ACC,BODY,True)
for i,(v,t) in enumerate([('2.24','Atualização passo a passo ao abrir'),('2.21','Legenda ao vivo e audiodescrição em todos os planos'),('2.20','Tradução ao vivo e audiodescrição com a própria voz'),('2.19','Audiodescrição no editor'),('2.18','Legenda ao vivo para transmissões'),('2.17','Mac e Linux'),('2.16','Status da conexão e diagnóstico'),('2.15','Formatos de entrega, prazos, histórico de versões')]):
    y=1.35+i*0.68
    c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(7.4),Inches(y+0.1),Inches(0.2),Inches(0.2)); c.fill.solid(); c.fill.fore_color.rgb=BG; c.line.color.rgb=ACC; c.line.width=Pt(2)
    rich(s,7.8,y,5,0.6,[(v+'  ',ACC,BODY),(t,INK,BODY)],14)
footer(s,11)

# 12 planos
s=new('Apresente os planos. Todas as funções em todos os planos; muda o número de computadores. Valores sujeitos a alteração: confirme no site antes de enviar propostas.')
title(s,'Planos','Todas as funções em todos os planos.')
planos=[('Individual','R$ 49,90/mês','Para quem trabalha sozinho',['1 computador','Todas as funções','Atualizações incluídas'],False),('Teamwork','R$ 49,90/mês','+ R$ 29,90 por computador adicional',['2 ou mais computadores','Todas as funções','Área do cliente com licenças e produção'],True),('Business / Enterprise','Sob cotação','10 ou mais computadores',['Workspace da empresa','Perfis e clientes compartilhados','Produção da equipe'],False)]
for i,(n_,p_,sub,its,hl) in enumerate(planos):
    x=0.7+i*4.1; box(s,x,1.9,3.85,3.9,RGBColor(0x2B,0x26,0x1C) if hl else PANEL,ACC if hl else LINE)
    txt(s,x+0.3,2.1,3.3,0.4,n_,18,INK,BODY,True); txt(s,x+0.3,2.55,3.3,0.7,p_,26 if i<2 else 24,ACC,HEAD); txt(s,x+0.3,3.3,3.3,0.4,sub,12,MUT)
    for j,t in enumerate(its): rich(s,x+0.3,3.85+j*0.5,3.3,0.4,[('✓  ',ACC,BODY),(t,INK,BODY)],13)
box(s,0.7,6.0,12,0.85); txt(s,0.95,6.07,11.5,0.7,'Experimente antes: 24 horas de avaliação com tudo liberado; depois, modo grátis para criar projetos e traduzir à mão. A IA usa a chave do próprio cliente, sem margem. Valores mensais, sujeitos a alteração.',12.5,MUT,anchor=MSO_ANCHOR.MIDDLE)
footer(s,12)

# 13 cta
s=new('Encerramento: convide para a avaliação gratuita e ofereça uma demonstração com um arquivo do fluxo do cliente.')
logo(s); txt(s,0.7,2.3,6,0.3,'VAMOS CONVERSAR',12,ACC,BODY,True)
txt(s,0.7,2.7,6.2,2.2,['Teste o CNT','Legendas hoje.'],50,INK,HEAD,spacing=0.95)
txt(s,0.7,4.6,5.8,1.3,'Baixe o instalador, use tudo por 24 horas e veja quanto tempo sua equipe ganha por episódio. Se quiser, fazemos uma demonstração com um arquivo do seu fluxo.',16,MUT)
box(s,0.7,5.95,5.2,0.85,PANEL,ACC); txt(s,0.9,6.0,4.8,0.3,'SITE E DOWNLOAD',10,MUT,BODY,True); txt(s,0.9,6.27,4.8,0.5,'legendas.cntoficial.com.br',20,INK,BODY,True)
pic(s,'inicio.jpg',7.2,1.0,5.6); pic(s,'novo-trabalho.jpg',7.9,3.9,5.0)
prs.save('../CNT-Legendas-Apresentacao.pptx'); print('ok',len(prs.slides))
